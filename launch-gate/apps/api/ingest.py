"""CSV and XLSX upload parsing, validation, and malformed-row isolation."""

from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import math
import re
from dataclasses import dataclass, field
from pathlib import PurePath
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.api.config import NORMALIZED_COLUMN_ALIASES, normalize_column_label
from packages.rules.platform_normalization import normalize_platform


MAX_CSV_ROWS = 200
MAX_REJECTED_CELL_CHARACTERS = 500


class UploadValidationError(ValueError):
    """Raised when an entire upload cannot be processed safely."""

    def __init__(self, message: str, *, reason_code: str = "invalid_upload") -> None:
        super().__init__(message)
        self.reason_code = reason_code
        self.reason = message


@dataclass(frozen=True)
class RejectedRow:
    """A rejected source row retained as inert, bounded text for batch storage."""

    data_row_index: int
    row_number: object | None
    reason_code: str
    reason: str
    raw_cells: tuple[str, ...]


@dataclass(frozen=True)
class ParsedCsvUpload:
    """Valid normalized ads and isolated rejected rows from one upload."""

    columns: tuple[str, ...]
    rows: list[dict[str, str]]
    rejected_rows: list[RejectedRow] = field(default_factory=list)
    valid_row_indices: tuple[int, ...] = ()
    unmapped_columns: tuple[str, ...] = ()

    @property
    def row_hashes(self) -> tuple[str, ...]:
        """Stable hashes of normalized ad fields, excluding source row numbers."""
        return tuple(hash_normalized_row(row) for row in self.rows)

    @property
    def sheet_hash(self) -> str:
        """Stable hash of valid normalized rows in their source order."""
        return hash_normalized_sheet(self.rows)


def _hash_payload(value: object) -> str:
    encoded = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def hash_normalized_row(row: dict[str, str]) -> str:
    """Hash canonical ad content for approval checks; hierarchy is bound by the sheet hash."""
    canonical = _canonical_hash_fields(row)
    return _hash_payload({key: canonical[key] for key in ("platform", "text", "headline", "link", "creative")})


def hash_normalized_sheet(rows: list[dict[str, str]]) -> str:
    """Hash canonical fields from valid normalized rows in source order."""
    canonical = [_canonical_hash_fields(row) for row in rows]
    return _hash_payload(canonical)


def _canonical_hash_fields(row: dict[str, str]) -> dict[str, str]:
    raw_platform = row.get("platform", "")
    normalized_platform = normalize_platform(raw_platform)
    canonical = {
        "platform": normalized_platform.canonical_value or raw_platform,
        "text": row.get("ad_text", "") or row.get("body_copy", ""),
        "headline": row.get("headline", ""),
        "link": row.get("destination_url", ""),
        "creative": row.get("creative_file", "") or row.get("creative_url", ""),
    }
    hierarchy_fields = (
        "campaign_key", "ad_set_key", "ad_key", "campaign_name", "campaign_objective",
        "campaign_budget_amount", "campaign_special_category", "campaign_budget_strategy",
        "ad_set_name", "targeting_spec", "optimization_goal", "billing_event",
        "ad_set_budget_amount", "ad_name", "page_id", "instagram_actor_id", "placements", "schedule",
    )
    hierarchy = {key: row[key] for key in hierarchy_fields if key in row and row[key] != ""}
    if hierarchy:
        canonical["hierarchy"] = hierarchy
    return canonical


def _normalize_header(header: str) -> str:
    if not isinstance(header, str):
        raise UploadValidationError(
            "column names must be text", reason_code="invalid_header"
        )
    stripped = header.strip()
    key = normalize_column_label(stripped)
    if not key:
        raise UploadValidationError(
            "file contains an empty column name", reason_code="invalid_header"
        )
    alias = NORMALIZED_COLUMN_ALIASES.get(key)
    if alias:
        return alias
    return re.sub(r"[^a-z0-9]+", "_", stripped.casefold()).strip("_")


def _validate_columns(headers: list[str]) -> tuple[str, ...]:
    columns = tuple(_normalize_header(header) for header in headers)
    if len(set(columns)) != len(columns):
        raise UploadValidationError(
            "file contains duplicate normalized columns",
            reason_code="duplicate_columns",
        )

    missing: list[str] = []
    if "platform" not in columns:
        missing.append("platform")
    if "destination_url" not in columns:
        missing.append("link")
    if not {"ad_text", "headline"}.intersection(columns):
        missing.append("text or headline")
    if missing:
        raise UploadValidationError(
            "missing required columns: " + ", ".join(missing),
            reason_code="missing_required_columns",
        )
    return columns


def _unmapped_column_headers(headers: list[str]) -> tuple[str, ...]:
    """Return original header names that did not resolve to a known alias."""
    return tuple(
        header
        for header in headers
        if normalize_column_label(header.strip()) not in NORMALIZED_COLUMN_ALIASES
    )


def _is_empty(value: object | None) -> bool:
    return value is None or str(value).strip() == ""


def _is_empty_record(record: list[object] | tuple[object, ...]) -> bool:
    return all(_is_empty(value) for value in record)


def _neutralize_formula_cell(value: str) -> str:
    """Prefix formula-like cell text so spreadsheet software treats it as text."""
    if value.lstrip()[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _raw_cell_text(value: object | None) -> str:
    text = "" if value is None else str(value)
    return _neutralize_formula_cell(text)[:MAX_REJECTED_CELL_CHARACTERS]


def _parse_positive_row_number(value: object | None) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return None
        parsed = int(value)
    elif isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return None
        try:
            parsed = int(stripped)
        except ValueError:
            try:
                decimal_value = Decimal(stripped)
            except InvalidOperation:
                return None
            if not decimal_value.is_finite() or decimal_value != decimal_value.to_integral_value():
                return None
            parsed = int(decimal_value)
    else:
        return None
    return parsed if parsed > 0 else None


def _output_columns(source_columns: tuple[str, ...]) -> tuple[str, ...]:
    if "row_number" in source_columns:
        return source_columns
    return ("row_number", *source_columns)


def _make_rejected_row(
    *,
    data_row_index: int,
    source_columns: tuple[str, ...],
    raw_values: list[object] | tuple[object, ...],
    reason_code: str,
    reason: str,
) -> RejectedRow:
    row_number: object | None = None
    if "row_number" in source_columns:
        position = source_columns.index("row_number")
        if position < len(raw_values):
            row_number = raw_values[position]
    return RejectedRow(
        data_row_index=data_row_index,
        row_number=row_number,
        reason_code=reason_code,
        reason=reason,
        raw_cells=tuple(_raw_cell_text(value) for value in raw_values),
    )


def _validate_record(
    *,
    source_columns: tuple[str, ...],
    record: list[object] | tuple[object, ...],
    data_row_index: int,
    width_mismatch: bool,
) -> tuple[dict[str, str] | None, RejectedRow | None]:
    if _is_empty_record(record):
        return None, _make_rejected_row(
            data_row_index=data_row_index,
            source_columns=source_columns,
            raw_values=record,
            reason_code="empty_record",
            reason="record is completely empty",
        )

    if width_mismatch:
        return None, _make_rejected_row(
            data_row_index=data_row_index,
            source_columns=source_columns,
            raw_values=record,
            reason_code="wrong_width",
            reason="record field count differs from the header column count",
        )

    field_values = list(record[: len(source_columns)])
    field_values.extend([None] * (len(source_columns) - len(field_values)))
    raw_by_column = dict(zip(source_columns, field_values))
    row_number: int
    if "row_number" in source_columns:
        parsed_row_number = _parse_positive_row_number(raw_by_column.get("row_number"))
        if parsed_row_number is None:
            return None, _make_rejected_row(
                data_row_index=data_row_index,
                source_columns=source_columns,
                raw_values=record,
                reason_code="invalid_row_number",
                reason="row number must be a positive integer or integral float",
            )
        row_number = parsed_row_number
    else:
        row_number = data_row_index

    invalid_fields: list[str] = []
    if _is_empty(raw_by_column.get("platform")):
        invalid_fields.append("platform")
    if _is_empty(raw_by_column.get("destination_url")):
        invalid_fields.append("link")
    if _is_empty(raw_by_column.get("ad_text")) and _is_empty(
        raw_by_column.get("headline")
    ):
        invalid_fields.append("text or headline")
    if invalid_fields:
        field_text = ", ".join(invalid_fields)
        return None, _make_rejected_row(
            data_row_index=data_row_index,
            source_columns=source_columns,
            raw_values=record,
            reason_code="missing_required_values",
            reason="required values are empty: " + field_text,
        )

    normalized = {
        column: _neutralize_formula_cell("" if value is None else str(value))
        for column, value in zip(source_columns, field_values)
    }
    raw_ad_text = normalized.get("ad_text", "")
    headline = normalized.get("headline", "")
    if _is_empty(headline) and " / " in raw_ad_text:
        ad_text, parsed_headline = raw_ad_text.rsplit(" / ", 1)
        if not _is_empty(ad_text) and not _is_empty(parsed_headline):
            normalized["raw_ad_text"] = raw_ad_text
            normalized["ad_text"] = ad_text
            normalized["headline"] = parsed_headline
    normalized["row_number"] = str(row_number)
    return normalized, None


def _validate_filename_and_size(
    filename: str,
    content: bytes,
    *,
    extension: str,
    max_file_bytes: int,
    unsupported_type_message: str | None = None,
) -> None:
    if not isinstance(filename, str) or not filename.strip():
        raise UploadValidationError("a filename is required", reason_code="invalid_filename")
    if PurePath(filename).suffix.casefold() != extension:
        raise UploadValidationError(
            unsupported_type_message or f"only {extension[1:].upper()} files are accepted",
            reason_code="unsupported_file_type",
        )
    if not isinstance(content, bytes):
        raise UploadValidationError("upload content must be bytes", reason_code="invalid_content")
    if (
        not isinstance(max_file_bytes, int)
        or isinstance(max_file_bytes, bool)
        or max_file_bytes <= 0
    ):
        raise UploadValidationError(
            "max_file_bytes must be a positive integer", reason_code="invalid_size_limit"
        )
    if len(content) > max_file_bytes:
        raise UploadValidationError(
            "upload exceeds the configured byte limit", reason_code="file_too_large"
        )


def _finish_upload(
    *,
    source_columns: tuple[str, ...],
    records: list[list[object] | tuple[object, ...]],
    valid_widths: list[bool],
    unmapped_columns: tuple[str, ...] = (),
) -> ParsedCsvUpload:
    # A trailing fully empty record is treated like trailing whitespace in the file.
    while records and _is_empty_record(records[-1]):
        records.pop()
        valid_widths.pop()

    rows: list[dict[str, str]] = []
    valid_row_indices: list[int] = []
    rejected_rows: list[RejectedRow] = []
    for data_row_index, (record, width_ok) in enumerate(
        zip(records, valid_widths), start=1
    ):
        normalized, rejected = _validate_record(
            source_columns=source_columns,
            record=record,
            data_row_index=data_row_index,
            width_mismatch=not width_ok,
        )
        if rejected is not None:
            rejected_rows.append(rejected)
            continue
        assert normalized is not None
        rows.append(normalized)
        valid_row_indices.append(data_row_index)
        if len(rows) > MAX_CSV_ROWS:
            raise UploadValidationError(
                f"upload may contain at most {MAX_CSV_ROWS} valid ads",
                reason_code="too_many_ads",
            )

    return ParsedCsvUpload(
        columns=_output_columns(source_columns),
        rows=rows,
        rejected_rows=rejected_rows,
        valid_row_indices=tuple(valid_row_indices),
        unmapped_columns=unmapped_columns,
    )


def parse_csv_upload(
    filename: str,
    content: bytes,
    *,
    max_file_bytes: int,
) -> ParsedCsvUpload:
    """Validate a CSV upload, isolating malformed rows and preserving valid rows."""
    _validate_filename_and_size(
        filename, content, extension=".csv", max_file_bytes=max_file_bytes
    )
    if b"\x00" in content:
        raise UploadValidationError(
            "CSV contains a NUL byte", reason_code="unparseable_csv"
        )
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UploadValidationError(
            "CSV must be UTF-8 encoded", reason_code="unparseable_csv"
        ) from exc

    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    try:
        headers = next(reader, None)
        if not headers:
            raise UploadValidationError(
                "missing required columns: platform, link, text or headline",
                reason_code="missing_required_columns",
            )
        source_columns = _validate_columns(headers)
        unmapped_columns = _unmapped_column_headers(headers)
        records: list[list[str]] = list(reader)
    except csv.Error as exc:
        raise UploadValidationError(
            f"CSV is unparseable: {exc}", reason_code="unparseable_csv"
        ) from exc

    return _finish_upload(
        source_columns=source_columns,
        records=records,
        valid_widths=[len(record) == len(source_columns) for record in records],
        unmapped_columns=unmapped_columns,
    )


def parse_xlsx_upload(
    filename: str,
    content: bytes,
    *,
    max_file_bytes: int,
    max_uncompressed_bytes: int | None = None,
    max_archive_entries: int = 200,
) -> ParsedCsvUpload:
    """Validate and isolate rows from the sole non-empty worksheet in an XLSX."""
    if isinstance(filename, str) and PurePath(filename).suffix.casefold() == ".xlsm":
        _validate_filename_and_size(
            filename, content, extension=".xlsm", max_file_bytes=max_file_bytes
        )
        raise UploadValidationError(
            "This is a macro-enabled workbook; macro-enabled XLSX is not accepted. Upload a non-macro XLSX file.",
            reason_code="macro_enabled_xlsx",
        )
    _validate_filename_and_size(
        filename,
        content,
        extension=".xlsx",
        max_file_bytes=max_file_bytes,
        unsupported_type_message="Please upload a valid XLSX file.",
    )
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            archive_entries = archive.infolist()
            if len(archive_entries) > max_archive_entries:
                raise UploadValidationError(
                    f"XLSX archive has more than {max_archive_entries} entries",
                    reason_code="xlsx_too_many_entries",
                )
            uncompressed_size = sum(entry.file_size for entry in archive_entries)
            if (
                max_uncompressed_bytes is not None
                and uncompressed_size > max_uncompressed_bytes
            ):
                raise UploadValidationError(
                    "XLSX uncompressed archive size "
                    f"{uncompressed_size} bytes exceeds the configured limit of "
                    f"{max_uncompressed_bytes} bytes",
                    reason_code="xlsx_uncompressed_too_large",
                )
            names = [entry.filename for entry in archive_entries]
            if any(name.casefold().endswith("vbaproject.bin") for name in names):
                raise UploadValidationError(
                    "This is a macro-enabled workbook; macro-enabled XLSX is not accepted. Upload a non-macro XLSX file.",
                    reason_code="macro_enabled_xlsx",
                )
            try:
                content_types = archive.read("[Content_Types].xml").decode("utf-8")
            except (KeyError, UnicodeDecodeError) as exc:
                raise UploadValidationError(
                    "XLSX is corrupt or unreadable; upload a valid XLSX file.",
                    reason_code="unreadable_xlsx",
                ) from exc
            if "macroenabled" in content_types.casefold():
                raise UploadValidationError(
                    "This is a macro-enabled workbook; macro-enabled XLSX is not accepted. Upload a non-macro XLSX file.",
                    reason_code="macro_enabled_xlsx",
                )
    except UploadValidationError:
        raise
    except Exception as exc:
        raise UploadValidationError(
            "XLSX is corrupt, encrypted, or unreadable; upload a valid XLSX file.",
            reason_code="unreadable_xlsx",
        ) from exc

    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except (InvalidFileException, OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise UploadValidationError(
            "XLSX is corrupt, encrypted, or unreadable; upload a valid XLSX file.",
            reason_code="unreadable_xlsx",
        ) from exc

    try:
        non_empty_sheets: list[tuple[str, list[tuple[object, ...]]]] = []
        for worksheet in workbook.worksheets:
            worksheet_rows = [
                tuple(record)
                for record in worksheet.iter_rows(values_only=True)
            ]
            if any(
                not _is_empty(value)
                for record in worksheet_rows
                for value in record
            ):
                non_empty_sheets.append((worksheet.title, worksheet_rows))

        if len(non_empty_sheets) > 1:
            raise UploadValidationError(
                "XLSX must contain exactly one non-empty worksheet.",
                reason_code="multiple_worksheets",
            )
        if not non_empty_sheets:
            raise UploadValidationError(
                "XLSX has no non-empty worksheet; upload a valid XLSX file with a worksheet.",
                reason_code="unreadable_xlsx",
            )

        _, worksheet_rows = non_empty_sheets[0]
        header_values = worksheet_rows[0] if worksheet_rows else None
        if not header_values:
            raise UploadValidationError(
                "missing required columns: platform, link, text or headline",
                reason_code="missing_required_columns",
            )
        header_values = list(header_values)
        while header_values and _is_empty(header_values[-1]):
            header_values.pop()
        if any(_is_empty(value) for value in header_values):
            raise UploadValidationError(
                "XLSX contains an empty column name", reason_code="invalid_header"
            )
        raw_headers = [str(value) for value in header_values]
        source_columns = _validate_columns(raw_headers)
        unmapped_columns = _unmapped_column_headers(raw_headers)

        records = worksheet_rows[1:]
        width_checks = [
            not any(not _is_empty(value) for value in record[len(source_columns):])
            for record in records
        ]
    except UploadValidationError:
        raise
    except Exception as exc:
        raise UploadValidationError(
            "XLSX is corrupt, encrypted, or unreadable; upload a valid XLSX file.",
            reason_code="unreadable_xlsx",
        ) from exc
    finally:
        workbook.close()

    return _finish_upload(
        source_columns=source_columns,
        records=records,
        valid_widths=width_checks,
        unmapped_columns=unmapped_columns,
    )
