from io import BytesIO
import zipfile

import pytest
from openpyxl import Workbook

from apps.api.ingest import (
    MAX_CSV_ROWS,
    ParsedCsvUpload,
    UploadValidationError,
    parse_xlsx_upload,
)


MAX_BYTES = 100_000


def make_workbook(*, extra_sheet: str | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ads"
    sheet.append(["#", "Channel", "Primary Text", "Headline", "URL", "Asset"])
    sheet.append([1, "Search", "hello", "title", "/page", None])
    if extra_sheet:
        other_sheet = workbook.create_sheet(extra_sheet)
        other_sheet.append(["Platform", "Headline", "URL"])
        other_sheet.append(["Meta", "another ad", "/other-page"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse(content: bytes, *, limit: int = MAX_BYTES):
    return parse_xlsx_upload(
        "ads.xlsx", content, max_file_bytes=limit
    )


def test_normalizes_xlsx_headers_and_matches_csv_result_shape():
    result = parse(make_workbook())

    assert isinstance(result, ParsedCsvUpload)
    assert result.columns == (
        "row_number",
        "platform",
        "ad_text",
        "headline",
        "destination_url",
        "creative_file",
    )
    assert result.rows == [
        {
            "row_number": "1",
            "platform": "Search",
            "ad_text": "hello",
            "headline": "title",
            "destination_url": "/page",
            "creative_file": "",
        }
    ]


def test_rejects_multiple_non_empty_worksheets():
    content = make_workbook(extra_sheet="Other")
    with pytest.raises(UploadValidationError) as exc_info:
        parse(content)
    assert exc_info.value.reason_code == "multiple_worksheets"


def test_ignores_empty_worksheets_when_exactly_one_contains_data():
    workbook = Workbook()
    first = workbook.active
    first.title = "Ads"
    first.append(["Row", "Platform", "Headline", "URL"])
    first.append([1, "Meta", "hello", "/page"])
    workbook.create_sheet("Empty")
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    assert len(parse(output.getvalue()).rows) == 1


def test_formula_cells_are_neutralized():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ads"
    sheet.append(["Row", "Platform", "Headline", "URL"])
    sheet.append([1, "Meta", "=1+1", "/page"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    result = parse(output.getvalue())
    assert result.rows[0]["headline"] == "'=1+1"


def test_rejects_macro_enabled_content_and_non_xlsx_extensions():
    original = make_workbook()
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(original)) as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("xl/vbaProject.bin", b"synthetic macro marker")

    with pytest.raises(UploadValidationError, match="macro-enabled"):
        parse(output.getvalue())
    with pytest.raises(UploadValidationError, match="non-macro XLSX"):
        parse_xlsx_upload(
            "ads.xlsm", original, max_file_bytes=MAX_BYTES
        )


def test_validates_size_workbook_format_and_headers():
    content = make_workbook()
    with pytest.raises(UploadValidationError, match="byte limit"):
        parse(content, limit=len(content) - 1)
    with pytest.raises(UploadValidationError, match="valid XLSX"):
        parse(b"not a workbook")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ads"
    sheet.append(["Platform", "Headline"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(
        UploadValidationError, match="missing required columns: link"
    ) as exc_info:
        parse(output.getvalue())
    assert exc_info.value.reason_code == "missing_required_columns"


def test_enforces_the_200_data_row_limit():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ads"
    sheet.append(["Row", "Platform", "Headline", "URL"])
    for row_number in range(1, MAX_CSV_ROWS + 2):
        sheet.append([row_number, "Meta", "Hello", "/page"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(UploadValidationError, match="at most 200"):
        parse(output.getvalue())
