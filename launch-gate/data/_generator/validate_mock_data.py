import csv, json, re, sys, hashlib
from pathlib import Path
from urllib.parse import urlparse, parse_qsl
from openpyxl import load_workbook
D = Path(sys.argv[1]); L = json.loads((D/"platform_limits.json").read_text())["platforms"]
N = json.loads((D/"naming_convention.json").read_text()); RI = json.loads((D/"clients/rule_index.json").read_text())
MAN = {r["file"]:(int(r["width"]),int(r["height"])) for r in csv.DictReader(open(D/"creatives_manifest.csv"))}
RAT={"1:1":1.0,"4:5":.8,"9:16":.5625,"1.91:1":1.91}
def rok(w,h,al): return any(abs(w/h-RAT[a])/RAT[a]<.02 for a in al)
bad=[]; checked=0
for i in range(1,31):
    lab=json.loads((D/f"labels/sheet_{i:02d}.labels.json").read_text())
    f=D/"labeled_sheets"/lab["sheet"]
    assert hashlib.sha256(f.read_bytes()).hexdigest()==lab["sheet_sha256"]
    if f.suffix==".csv": rows=list(csv.reader(open(f,encoding="utf-8")))[1:]
    else: rows=[["" if c is None else c for c in r] for r in load_workbook(f).active.iter_rows(values_only=True)][1:]
    assert len(rows)==len(lab["rows"])==lab["data_rows"], (f, len(rows), len(lab["rows"]))
    assert 40<=len(rows)<=204
    for cells,l in zip(rows,lab["rows"]):
        assert str(cells[0])==str(l["row_id"])
        if l["malformed"]: continue
        t=l["truth"]; pl=l["platform"]; errs={(e["type"],e["field"]) for e in l["planted_errors"]}
        for e in l["planted_errors"]:
            for k in ("rule_ref",):
                r=e[k]
                if r and r[:2] in ("BL","CL","OH") and "NAMING" not in r:
                    c={"BL":"brightleaf","CL":"calmleaf","OH":"oakandhome"}[r[:2]]
                    assert r in RI[c]["register"]+RI[c]["brand_book"], (r)
                    assert l["client"]==c, ("cross-client ref", r, l["client"])
        lim=L[pl]; checked+=1
        tl=len(t["text"])>lim["text_max"]["value"]; hl=len(t["headline"])>lim["headline_max"]["value"]
        assert tl==(("text_too_long","text") in errs), (f.name,l["row_id"],"text",len(t["text"]))
        assert hl==(("headline_too_long","headline") in errs), (f.name,l["row_id"],"headline",len(t["headline"]))
        cr=t["creative"]
        if pl!="google":
            if cr=="": assert ("creative_missing","creative") in errs
            else:
                w,h=MAN[cr]; assert (not rok(w,h,lim["ratios"]["value"]))==(("creative_ratio_mismatch","creative") in errs),(f.name,l["row_id"],cr)
        if l["client"] is None: continue
        u=urlparse(t["link"] if "://" in t["link"] else "//"+t["link"])
        q=u.query; mech_utm={e["type"] for e in l["planted_errors"] if e["type"].startswith("utm")}
        exp={"utm_source":pl,"utm_medium":N["medium"][pl]}
        params=dict(parse_qsl(q.lstrip("?"),keep_blank_values=True))
        canon=all(params.get(k)==v for k,v in exp.items()) and bool(re.match(N["campaign_pattern"] if l["client"]!="oakandhome" else N["client_overrides"]["oakandhome"]["campaign_pattern"], params.get("utm_campaign","")))
        if canon and not q.startswith("?") and "utm_sorce" not in q: assert not mech_utm,(f.name,l["row_id"],"clean utm but labelled",mech_utm,t["link"])
        else: assert mech_utm,(f.name,l["row_id"],"bad utm but not labelled",t["link"])
        if l["expected_verdict"]=="ready":
            assert not l["planted_errors"]
print("sheets ok; ad rows checked:",checked)
a=list(csv.DictReader(open(D/"adversarial_rows.csv",encoding="utf-8"))); assert len(a)==12
ap=list(csv.DictReader(open(D/"approval_history.csv"))); print("approvals",len(ap),"mismatch",sum(r["hashes_match"]=="NO" for r in ap))
print("ALL CHECKS PASSED")
